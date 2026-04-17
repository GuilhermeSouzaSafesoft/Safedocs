import base64
import traceback
import zipfile
from io import BytesIO

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from docx import Document

from api.schemas import (
    AppendHistoryTableRequest,
    DocumentSchema,
    HealthResponse,
    PowerAutomateResponse,
    RootResponse,
)
from api.services import generate_docx_from_payload
from safedocs.blocks import render_block
from safedocs.models import InvalidDocumentError

app = FastAPI(
    title="Safedocs Docx API",
    version="0.1.0",
    servers=[
        {"url": "https://safedocs.onrender.com", "description": "Production"}
    ],
)


@app.get("/", response_model=RootResponse)
def root() -> RootResponse:
    return RootResponse(status="ok", message="Safedocs API pronta para gerar DOCX")


@app.get("/health", response_model=HealthResponse)
def health() -> HealthResponse:
    return HealthResponse(status="healthy")


def _payload_to_dict(payload: DocumentSchema) -> dict:
    if hasattr(payload, "model_dump"):
        return payload.model_dump(exclude_none=True)
    return payload.dict(exclude_none=True)


def _update_existing_revision_paragraph(document: Document, revision: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Revisão: "):
            paragraph.text = f"Revisão: {revision}"
            return


@app.post("/generate-docx")
def generate_docx(payload: DocumentSchema) -> FileResponse:
    try:
        doc_path = generate_docx_from_payload(_payload_to_dict(payload))
    except InvalidDocumentError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    return FileResponse(
        path=str(doc_path),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=doc_path.name,
        headers={"Content-Disposition": f'attachment; filename="{doc_path.name}"'},
    )


@app.post("/generate-docx-action", response_model=PowerAutomateResponse)
def generate_docx_action(payload: DocumentSchema) -> PowerAutomateResponse:
    try:
        doc_path = generate_docx_from_payload(_payload_to_dict(payload))
    except InvalidDocumentError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception:
        raise HTTPException(status_code=500, detail="Erro interno ao gerar o DOCX.")

    with doc_path.open("rb") as f:
        encoded = base64.b64encode(f.read()).decode("ascii")

    return PowerAutomateResponse(
        success=True,
        filename=doc_path.name,
        content_base64=encoded,
    )


@app.post("/append-history-table", response_model=PowerAutomateResponse)
def append_history_table(payload: AppendHistoryTableRequest) -> PowerAutomateResponse:
    try:
        try:
            source_bytes = base64.b64decode(payload.content_base64)
        except Exception as exc:
            raise ValueError(f"Erro ao decodificar base64: {exc}")

        try:
            with zipfile.ZipFile(BytesIO(source_bytes), "r") as archive:
                names = archive.namelist()
                print("DEBUG - arquivos no DOCX:", names[:20])
                if (
                    "[Content_Types].xml" not in names
                    or "word/document.xml" not in names
                ):
                    raise ValueError("Arquivo recebido não é um DOCX válido.")
        except zipfile.BadZipFile:
            raise ValueError(
                "Conteúdo recebido não é um arquivo ZIP válido (DOCX corrompido)."
            )

        try:
            document = Document(BytesIO(source_bytes))
        except Exception as exc:
            raise ValueError(f"Erro ao abrir DOCX com python-docx: {exc}")

        if payload.historico:
            _update_existing_revision_paragraph(
                document=document, revision=str(payload.historico[-1].versao)
            )

        document.add_paragraph()
        document.add_paragraph("Histórico de Revisões")

        historico_block = {
            "type": "tabela",
            "columns": ["Versão", "Data", "Autor"],
            "rows": [
                [str(item.versao), str(item.data), str(item.autor)]
                for item in payload.historico
            ],
        }
        render_block(document, historico_block)

        output = BytesIO()
        document.save(output)
        encoded = base64.b64encode(output.getvalue()).decode("ascii")

        filename = payload.filename or "documento"
        if filename.lower().endswith(".docx"):
            filename = f"{filename[:-5]}_modificado.docx"
        else:
            filename = f"{filename}_modificado.docx"

        return PowerAutomateResponse(
            success=True,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            content_base64=encoded,
        )
    except Exception as exc:
        print("ERRO NO ENDPOINT /append-history-table:")
        print(repr(exc))
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(exc))
